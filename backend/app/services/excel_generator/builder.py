from io import BytesIO
from typing import Any
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.excel_generator.blueprints import (
    BLUEPRINTS,
    get_instructions_for_blueprint,
)


class WorkbookBuilder:
    """
    Ensambla libros de Excel con diseño ejecutivo corporativo usando openpyxl.
    Genera tanto el archivo para el candidato como el archivo resuelto para calificación.
    """

    # Paleta de Colores Ejecutiva
    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    
    ACCENT_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
    ACCENT_HEADER = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")  # Teal 700
    
    BANNER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
    BANNER_FONT = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    
    CARD_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # Slate 50
    CARD_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    
    DATA_FONT = Font(name="Segoe UI", size=9)
    BOLD_FONT = Font(name="Segoe UI", size=9, bold=True)
    MONO_FONT = Font(name="Consolas", size=9)
    
    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    
    TOTAL_BORDER = Border(
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="0F172A"),
    )

    @classmethod
    def build_candidate_and_solution_workbooks(
        cls,
        blueprint_id: str,
        title: str,
        data: dict[str, Any],
        difficulty: str = "intermediate",
    ) -> tuple[bytes, bytes, dict[str, Any]]:
        """
        Genera el par de libros en memoria: (candidate_bytes, solution_bytes, expected_summary).
        """
        bp = BLUEPRINTS.get(blueprint_id, BLUEPRINTS["pivot_sales_fulfillment"])

        if blueprint_id == "pivot_sales_fulfillment":
            return cls._build_pivot_fulfillment(bp, title, data, difficulty)
        elif blueprint_id == "vlookup_catalog_pricing":
            return cls._build_vlookup_pricing(bp, title, data, difficulty)
        elif blueprint_id == "logic_kpi_evaluation":
            return cls._build_logic_kpi(bp, title, data, difficulty)
        elif blueprint_id == "audit_reconciliation":
            return cls._build_audit_reconciliation(bp, title, data, difficulty)
        else:
            return cls._build_pivot_fulfillment(bp, title, data, difficulty)

    # ---------------------------------------------------------
    # 1. BLUEPRINT: Tablas Dinámicas y Surtimiento
    # ---------------------------------------------------------
    @classmethod
    def _build_pivot_fulfillment(
        cls, bp: Any, title: str, data: dict[str, Any], difficulty: str = "intermediate"
    ) -> tuple[bytes, bytes, dict[str, Any]]:
        c_wb = openpyxl.Workbook()
        c_ws_task = c_wb.active
        c_ws_task.title = bp.task_sheet_name
        c_ws_data = c_wb.create_sheet(title=bp.source_sheet_name)
        
        s_wb = openpyxl.Workbook()
        s_ws_task = s_wb.active
        s_ws_task.title = bp.task_sheet_name
        s_ws_data = s_wb.create_sheet(title=bp.source_sheet_name)

        rows = data["rows"]

        # 1. Fill BaseDatos in both workbooks
        headers = ["FOLIO", "FECHA", "MES", "SUCURSAL", "CLAVE_PRODUCTO", "DESCRIPCION", "CANTIDAD_SURTIDA", "PRECIO_UNITARIO", "IMPORTE_TOTAL"]
        for ws in (c_ws_data, s_ws_data):
            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = cls.HEADER_FILL
                cell.font = cls.HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            for row_idx, r in enumerate(rows, start=2):
                ws.append([
                    r["FOLIO"],
                    r["FECHA"],
                    r["MES"],
                    r["SUCURSAL"],
                    r["CLAVE_PRODUCTO"],
                    r["DESCRIPCION"],
                    r["CANTIDAD_SURTIDA"],
                    r["PRECIO_UNITARIO"],
                    r["IMPORTE_TOTAL"],
                ])
                ws.cell(row=row_idx, column=2).number_format = "yyyy-mm-dd"
                ws.cell(row=row_idx, column=7).number_format = "#,##0"
                ws.cell(row=row_idx, column=8).number_format = "$#,##0.00"
                ws.cell(row=row_idx, column=9).number_format = "$#,##0.00"
            cls._auto_fit_columns(ws)

        # 2. Setup Task Sheet Header & Instructions in both
        inst_text = get_instructions_for_blueprint(bp.id, difficulty)
        for ws in (c_ws_task, s_ws_task):
            cls._create_instruction_banner(ws, f"{title} ({difficulty.upper()})", inst_text)

        # 3. Setup Target Control Cells Area based on difficulty
        tot_importe = round(sum(r["IMPORTE_TOTAL"] for r in rows), 2)
        tot_piezas = sum(r["CANTIDAD_SURTIDA"] for r in rows)
        avg_importe = round(tot_importe / len(rows), 2) if rows else 0.0
        tot_ene = round(sum(r["IMPORTE_TOTAL"] for r in rows if r["MES"] == "Enero"), 2)
        tot_feb = round(sum(r["IMPORTE_TOTAL"] for r in rows if r["MES"] == "Febrero"), 2)

        if difficulty == "basic":
            metric_configs = [
                ("Total General de Importe Surtido", "B13", tot_importe, "$#,##0.00"),
                ("Total Importe en Enero", "B14", tot_ene, "$#,##0.00"),
                ("Total Importe en Febrero", "B15", tot_feb, "$#,##0.00"),
            ]
            criteria_mapping = {
                "Totales de Importe": ["B13", "B14", "B15"],
            }
        elif difficulty == "advanced":
            tot_mar = round(sum(r["IMPORTE_TOTAL"] for r in rows if r["MES"] == "Marzo"), 2)
            metric_configs = [
                ("Total General de Importe Surtido", "B13", tot_importe, "$#,##0.00"),
                ("Total General de Piezas Surtidas", "B14", tot_piezas, "#,##0"),
                ("Promedio de Importe por Orden", "B15", avg_importe, "$#,##0.00"),
                ("Total Importe en Enero", "B16", tot_ene, "$#,##0.00"),
                ("Total Importe en Febrero", "B17", tot_feb, "$#,##0.00"),
                ("Total Importe en Marzo", "B18", tot_mar, "$#,##0.00"),
            ]
            criteria_mapping = {
                "Totales de Importe": ["B13", "B16", "B17", "B18"],
                "Volumen de Piezas": ["B14"],
                "Promedio General": ["B15"],
            }
        else:  # intermediate
            metric_configs = [
                ("Total General de Importe Surtido", "B13", tot_importe, "$#,##0.00"),
                ("Total General de Piezas Surtidas", "B14", tot_piezas, "#,##0"),
                ("Promedio de Importe por Orden", "B15", avg_importe, "$#,##0.00"),
                ("Total Importe en Enero", "B16", tot_ene, "$#,##0.00"),
                ("Total Importe en Febrero", "B17", tot_feb, "$#,##0.00"),
            ]
            criteria_mapping = {
                "Totales de Importe": ["B13", "B16", "B17"],
                "Volumen de Piezas": ["B14"],
                "Promedio General": ["B15"],
            }

        for ws in (c_ws_task, s_ws_task):
            ws["A10"] = "RESUMEN DE CONTROL EJECUTIVO (Escribe tus totales calculados o conéctalos con fórmulas):"
            ws["A10"].font = cls.BOLD_FONT
            ws["A12"] = "Métrica / Indicador"
            ws["B12"] = "Valor Calculado"
            ws["A12"].fill = cls.ACCENT_HEADER
            ws["B12"].fill = cls.ACCENT_HEADER
            ws["A12"].font = cls.HEADER_FONT
            ws["B12"].font = cls.HEADER_FONT

        expected_summary: dict[str, Any] = {}
        for label, cell_coord, val, num_fmt in metric_configs:
            c_ws_task[f"A{cell_coord[1:]}"] = label
            c_ws_task[f"A{cell_coord[1:]}"].font = cls.DATA_FONT
            c_ws_task[f"A{cell_coord[1:]}"].border = cls.CARD_BORDER
            c_ws_task[cell_coord].border = cls.CARD_BORDER
            c_ws_task[cell_coord].fill = cls.ACCENT_FILL

            s_ws_task[f"A{cell_coord[1:]}"] = label
            s_ws_task[f"A{cell_coord[1:]}"].font = cls.DATA_FONT
            s_ws_task[f"A{cell_coord[1:]}"].border = cls.CARD_BORDER
            s_ws_task[cell_coord].border = cls.CARD_BORDER
            s_ws_task[cell_coord] = val
            s_ws_task[cell_coord].number_format = num_fmt
            expected_summary[cell_coord] = val

        pivot_row = len(metric_configs) + 14
        c_ws_task[f"A{pivot_row}"] = "ZONA PARA TU TABLA DINÁMICA (Inserta tu tabla dinámica debajo de esta línea):"
        c_ws_task[f"A{pivot_row}"].font = cls.BOLD_FONT

        cls._auto_fit_columns(c_ws_task)
        cls._auto_fit_columns(s_ws_task)

        # Build Criteria Sheet
        criterios_ws = s_wb.create_sheet(title="Criterios")
        criterios_ws.append(["Criterio", "Rango_Celdas"])
        for crit_name, cells in criteria_mapping.items():
            criterios_ws.append([crit_name, ",".join(cells)])

        expected_summary["__criteria__"] = criteria_mapping
        return cls._save_to_bytes(c_wb), cls._save_to_bytes(s_wb), expected_summary

    # ---------------------------------------------------------
    # 2. BLUEPRINT: Cruce de Catálogos con BUSCARV
    # ---------------------------------------------------------
    @classmethod
    def _build_vlookup_pricing(
        cls, bp: Any, title: str, data: dict[str, Any], difficulty: str = "intermediate"
    ) -> tuple[bytes, bytes, dict[str, Any]]:
        c_wb = openpyxl.Workbook()
        c_ws_task = c_wb.active
        c_ws_task.title = bp.task_sheet_name
        c_ws_trx = c_wb.create_sheet(title=bp.source_sheet_name)
        c_ws_cat = c_wb.create_sheet(title="Catalogos")
        
        s_wb = openpyxl.Workbook()
        s_ws_task = s_wb.active
        s_ws_task.title = bp.task_sheet_name
        s_ws_trx = s_wb.create_sheet(title=bp.source_sheet_name)
        s_ws_cat = s_wb.create_sheet(title="Catalogos")

        catalog = data["catalog"]
        transactions = data["transactions"]

        # 1. Fill Transacciones in both
        if difficulty == "advanced":
            trx_headers = ["FOLIO", "CODIGO_PRODUCTO", "CANTIDAD", "DESCUENTO_PCT"]
            for ws in (c_ws_trx, s_ws_trx):
                ws.append(trx_headers)
                for c in range(1, len(trx_headers) + 1):
                    ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                    ws.cell(row=1, column=c).font = cls.HEADER_FONT
                    ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center")
                for t in transactions:
                    ws.append([t["FOLIO"], t["CODIGO_PRODUCTO"], t["CANTIDAD"], t["DESCUENTO_PCT"]])
                    ws.cell(row=ws.max_row, column=3).number_format = "#,##0"
                    ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
                cls._auto_fit_columns(ws)
        else:
            trx_headers = ["FOLIO", "CODIGO_PRODUCTO", "CANTIDAD"]
            for ws in (c_ws_trx, s_ws_trx):
                ws.append(trx_headers)
                for c in range(1, len(trx_headers) + 1):
                    ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                    ws.cell(row=1, column=c).font = cls.HEADER_FONT
                    ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center")
                for t in transactions:
                    ws.append([t["FOLIO"], t["CODIGO_PRODUCTO"], t["CANTIDAD"]])
                    ws.cell(row=ws.max_row, column=3).number_format = "#,##0"
                cls._auto_fit_columns(ws)

        # 2. Fill Catalog in both
        if difficulty == "basic":
            cat_headers = ["CODIGO", "DESCRIPCION", "PRECIO_UNITARIO"]
            for ws in (c_ws_cat, s_ws_cat):
                ws.append(cat_headers)
                for c in range(1, len(cat_headers) + 1):
                    ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                    ws.cell(row=1, column=c).font = cls.HEADER_FONT
                for p in catalog:
                    ws.append([p["code"], p["name"], p["price"]])
                    ws.cell(row=ws.max_row, column=3).number_format = "$#,##0.00"
                cls._auto_fit_columns(ws)
        else:
            cat_headers = ["CODIGO", "DESCRIPCION", "CATEGORIA", "PRECIO_UNITARIO"]
            for ws in (c_ws_cat, s_ws_cat):
                ws.append(cat_headers)
                for c in range(1, len(cat_headers) + 1):
                    ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                    ws.cell(row=1, column=c).font = cls.HEADER_FONT
                for p in catalog:
                    ws.append([p["code"], p["name"], p["cat"], p["price"]])
                    ws.cell(row=ws.max_row, column=4).number_format = "$#,##0.00"
                cls._auto_fit_columns(ws)

        # 3. Instructions
        inst_text = get_instructions_for_blueprint(bp.id, difficulty)
        for ws in (c_ws_task, s_ws_task):
            cls._create_instruction_banner(ws, f"{title} ({difficulty.upper()})", inst_text)

        # 4. Transaction Matrix Headers in RealizaEjercicio based on difficulty
        if difficulty == "basic":
            task_headers = ["FOLIO", "CODIGO", "DESCRIPCION (Fórmula)", "PRECIO_UNIT (Fórmula)", "CANTIDAD", "TOTAL (Fórmula)"]
        elif difficulty == "advanced":
            task_headers = [
                "FOLIO", "CODIGO", "DESCRIPCION (Fórmula)", "CATEGORIA (Fórmula)", "PRECIO_UNIT (Fórmula)",
                "CANTIDAD", "DESCUENTO %", "DESC $ (Fórmula)", "SUBTOTAL NETO (Fórmula)", "IVA 16% (Fórmula)", "TOTAL FINAL (Fórmula)"
            ]
        else:  # intermediate
            task_headers = [
                "FOLIO", "CODIGO", "DESCRIPCION (Fórmula)", "CATEGORIA (Fórmula)", "PRECIO_UNIT (Fórmula)",
                "CANTIDAD", "SUBTOTAL (Fórmula)", "IVA 16% (Fórmula)", "TOTAL (Fórmula)"
            ]
        
        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=10, column=1, value="MATRIZ DE OPERACIONES A FORMULAR:").font = cls.BOLD_FONT
            for col_idx, h in enumerate(task_headers, start=1):
                cell = ws.cell(row=11, column=col_idx, value=h)
                cell.fill = cls.HEADER_FILL
                cell.font = cls.HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

        expected_summary = {}
        target_cells_lookup = []
        target_cells_math = []

        num_cols = len(task_headers)

        for idx, trx in enumerate(transactions, start=12):
            if difficulty == "basic":
                c_ws_task.append([trx["FOLIO"], trx["CODIGO_PRODUCTO"], "", "", trx["CANTIDAD"], ""])
                s_ws_task.append([
                    trx["FOLIO"],
                    trx["CODIGO_PRODUCTO"],
                    trx["_prod_name"],
                    trx["_unit_price"],
                    trx["CANTIDAD"],
                    trx["_subtotal"],
                ])
                s_ws_task.cell(row=idx, column=4).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "$#,##0.00"

                expected_summary[f"C{idx}"] = trx["_prod_name"]
                expected_summary[f"D{idx}"] = trx["_unit_price"]
                expected_summary[f"F{idx}"] = trx["_subtotal"]
                target_cells_lookup.extend([f"C{idx}", f"D{idx}"])
                target_cells_math.append(f"F{idx}")

            elif difficulty == "advanced":
                desc_pct = trx["DESCUENTO_PCT"]
                desc_amount = round(trx["_subtotal"] * desc_pct, 2)
                sub_neto = round(trx["_subtotal"] - desc_amount, 2)
                iva_neto = round(sub_neto * 0.16, 2)
                final_neto = round(sub_neto + iva_neto, 2)

                c_ws_task.append([trx["FOLIO"], trx["CODIGO_PRODUCTO"], "", "", "", trx["CANTIDAD"], desc_pct, "", "", "", ""])
                c_ws_task.cell(row=idx, column=7).number_format = "0.0%"

                s_ws_task.append([
                    trx["FOLIO"],
                    trx["CODIGO_PRODUCTO"],
                    trx["_prod_name"],
                    trx["_category"],
                    trx["_unit_price"],
                    trx["CANTIDAD"],
                    desc_pct,
                    desc_amount,
                    sub_neto,
                    iva_neto,
                    final_neto,
                ])
                s_ws_task.cell(row=idx, column=5).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=7).number_format = "0.0%"
                s_ws_task.cell(row=idx, column=8).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=9).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=10).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=11).number_format = "$#,##0.00"

                expected_summary[f"C{idx}"] = trx["_prod_name"]
                expected_summary[f"D{idx}"] = trx["_category"]
                expected_summary[f"E{idx}"] = trx["_unit_price"]
                expected_summary[f"H{idx}"] = desc_amount
                expected_summary[f"I{idx}"] = sub_neto
                expected_summary[f"K{idx}"] = final_neto

                target_cells_lookup.extend([f"C{idx}", f"D{idx}", f"E{idx}"])
                target_cells_math.extend([f"H{idx}", f"I{idx}", f"K{idx}"])

            else:  # intermediate
                c_ws_task.append([trx["FOLIO"], trx["CODIGO_PRODUCTO"], "", "", "", trx["CANTIDAD"], "", "", ""])
                subtot = round(trx["CANTIDAD"] * trx["_unit_price"], 2)
                iva = round(subtot * 0.16, 2)
                tot = round(subtot + iva, 2)
                s_ws_task.append([
                    trx["FOLIO"],
                    trx["CODIGO_PRODUCTO"],
                    trx["_prod_name"],
                    trx["_category"],
                    trx["_unit_price"],
                    trx["CANTIDAD"],
                    subtot,
                    iva,
                    tot,
                ])
                s_ws_task.cell(row=idx, column=5).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=7).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=8).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=9).number_format = "$#,##0.00"

                expected_summary[f"C{idx}"] = trx["_prod_name"]
                expected_summary[f"D{idx}"] = trx["_category"]
                expected_summary[f"E{idx}"] = trx["_unit_price"]
                expected_summary[f"G{idx}"] = subtot
                expected_summary[f"H{idx}"] = iva
                expected_summary[f"I{idx}"] = tot
                target_cells_lookup.extend([f"C{idx}", f"D{idx}", f"E{idx}"])
                target_cells_math.extend([f"G{idx}", f"H{idx}", f"I{idx}"])

            for col in range(1, num_cols + 1):
                c_ws_task.cell(row=idx, column=col).border = cls.THIN_BORDER
                s_ws_task.cell(row=idx, column=col).border = cls.THIN_BORDER

        # Add Totals Row
        tot_row = len(transactions) + 12
        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=tot_row, column=1, value="TOTAL GENERAL:").font = cls.BOLD_FONT
            for col in range(1, num_cols + 1):
                ws.cell(row=tot_row, column=col).border = cls.TOTAL_BORDER

        if difficulty == "basic":
            tot_final = round(sum(t["_subtotal"] for t in transactions), 2)
            s_ws_task.cell(row=tot_row, column=6, value=tot_final).number_format = "$#,##0.00"
            expected_summary[f"F{tot_row}"] = tot_final
            target_cells_math.append(f"F{tot_row}")
        elif difficulty == "advanced":
            tot_desc = round(sum(round(t["_subtotal"] * t["DESCUENTO_PCT"], 2) for t in transactions), 2)
            tot_sub_neto = round(sum(t["_subtotal"] - round(t["_subtotal"] * t["DESCUENTO_PCT"], 2) for t in transactions), 2)
            tot_final_neto = round(sum(round((t["_subtotal"] - round(t["_subtotal"] * t["DESCUENTO_PCT"], 2)) * 1.16, 2) for t in transactions), 2)
            s_ws_task.cell(row=tot_row, column=8, value=tot_desc).number_format = "$#,##0.00"
            s_ws_task.cell(row=tot_row, column=9, value=tot_sub_neto).number_format = "$#,##0.00"
            s_ws_task.cell(row=tot_row, column=11, value=tot_final_neto).number_format = "$#,##0.00"
            expected_summary[f"I{tot_row}"] = tot_sub_neto
            expected_summary[f"K{tot_row}"] = tot_final_neto
            target_cells_math.extend([f"I{tot_row}", f"K{tot_row}"])
        else:
            tot_subtotal = round(sum(round(t["CANTIDAD"] * t["_unit_price"], 2) for t in transactions), 2)
            tot_iva = round(sum(round(round(t["CANTIDAD"] * t["_unit_price"], 2) * 0.16, 2) for t in transactions), 2)
            tot_final = round(tot_subtotal + tot_iva, 2)
            s_ws_task.cell(row=tot_row, column=7, value=tot_subtotal).number_format = "$#,##0.00"
            s_ws_task.cell(row=tot_row, column=8, value=tot_iva).number_format = "$#,##0.00"
            s_ws_task.cell(row=tot_row, column=9, value=tot_final).number_format = "$#,##0.00"
            expected_summary[f"G{tot_row}"] = tot_subtotal
            expected_summary[f"H{tot_row}"] = tot_iva
            expected_summary[f"I{tot_row}"] = tot_final
            target_cells_math.extend([f"G{tot_row}", f"H{tot_row}", f"I{tot_row}"])

        expected_summary["__criteria__"] = {
            "Búsqueda de Catálogo": target_cells_lookup,
            "Cálculo de Precios e Impuestos": target_cells_math,
        }

        # Criteria sheet
        crit_ws = s_wb.create_sheet(title="Criterios")
        crit_ws.append(["Criterio", "Rango_Celdas"])
        crit_ws.append(["Búsqueda de Catálogo", f"C12:E{tot_row-1}"])
        crit_ws.append(["Cálculo de Precios e Impuestos", f"F12:K{tot_row}"])

        cls._auto_fit_columns(c_ws_task)
        cls._auto_fit_columns(s_ws_task)

        return cls._save_to_bytes(c_wb), cls._save_to_bytes(s_wb), expected_summary

    # ---------------------------------------------------------
    # 3. BLUEPRINT: Lógica y KPIs
    # ---------------------------------------------------------
    @classmethod
    def _build_logic_kpi(
        cls, bp: Any, title: str, data: dict[str, Any], difficulty: str = "basic"
    ) -> tuple[bytes, bytes, dict[str, Any]]:
        c_wb = openpyxl.Workbook()
        c_ws_task = c_wb.active
        c_ws_task.title = bp.task_sheet_name
        c_ws_src = c_wb.create_sheet(title=bp.source_sheet_name)
        
        s_wb = openpyxl.Workbook()
        s_ws_task = s_wb.active
        s_ws_task.title = bp.task_sheet_name
        s_ws_src = s_wb.create_sheet(title=bp.source_sheet_name)

        rows = data["rows"]

        # 1. Fill RegistroMetas in both
        src_headers = ["ID_EMPLEADO", "NOMBRE", "SUCURSAL", "META_ASIGNADA", "LOGRO_REAL"]
        for ws in (c_ws_src, s_ws_src):
            ws.append(src_headers)
            for c in range(1, len(src_headers) + 1):
                ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                ws.cell(row=1, column=c).font = cls.HEADER_FONT
                ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center")
            for r in rows:
                ws.append([r["ID_EMPLEADO"], r["NOMBRE"], r["SUCURSAL"], r["META_ASIGNADA"], r["LOGRO_REAL"]])
                ws.cell(row=ws.max_row, column=4).number_format = "$#,##0.00"
                ws.cell(row=ws.max_row, column=5).number_format = "$#,##0.00"
            cls._auto_fit_columns(ws)

        # 2. Instructions
        inst_text = get_instructions_for_blueprint(bp.id, difficulty)
        for ws in (c_ws_task, s_ws_task):
            cls._create_instruction_banner(ws, f"{title} ({difficulty.upper()})", inst_text)

        # 3. Task Headers based on difficulty
        if difficulty == "advanced":
            headers = [
                "ID_EMPLEADO", "NOMBRE", "SUCURSAL", "META_ASIGNADA", "LOGRO_REAL",
                "% CUMPLIMIENTO (Fórmula)", "ESTATUS (Fórmula)", "COMISION VAR (Fórmula)", "BONO FIJO (Fórmula)", "COMPENSACION TOTAL (Fórmula)"
            ]
        else:
            headers = [
                "ID_EMPLEADO", "NOMBRE", "SUCURSAL", "META_ASIGNADA", "LOGRO_REAL",
                "% CUMPLIMIENTO (Fórmula)", "ESTATUS (Fórmula)", "BONO (Fórmula)"
            ]

        num_cols = len(headers)

        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=10, column=1, value="EVALUACIÓN DE DESEMPEÑO POR COLABORADOR:").font = cls.BOLD_FONT
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=11, column=col_idx, value=h)
                cell.fill = cls.HEADER_FILL
                cell.font = cls.HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

        expected_summary = {}
        criteria_status = []
        criteria_bonus = []

        for idx, r in enumerate(rows, start=12):
            pct = round(r["LOGRO_REAL"] / max(r["META_ASIGNADA"], 1.0), 4)

            if difficulty == "basic":
                estatus = "CUMPLIDO" if pct >= 1.0 else "NO CUMPLIDO"
                bono = 3000.0 if estatus == "CUMPLIDO" else 0.0

                c_ws_task.append([r["ID_EMPLEADO"], r["NOMBRE"], r["SUCURSAL"], r["META_ASIGNADA"], r["LOGRO_REAL"], "", "", ""])
                s_ws_task.append([
                    r["ID_EMPLEADO"],
                    r["NOMBRE"],
                    r["SUCURSAL"],
                    r["META_ASIGNADA"],
                    r["LOGRO_REAL"],
                    pct,
                    estatus,
                    bono,
                ])
                s_ws_task.cell(row=idx, column=4).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=5).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "0.0%"
                s_ws_task.cell(row=idx, column=8).number_format = "$#,##0.00"

                expected_summary[f"G{idx}"] = estatus
                expected_summary[f"H{idx}"] = bono
                criteria_status.append(f"G{idx}")
                criteria_bonus.append(f"H{idx}")

            elif difficulty == "advanced":
                if pct >= 1.15:
                    estatus = "SOBRESALIENTE"
                    comision_pct = 0.10
                elif pct >= 1.0:
                    estatus = "CUMPLIDO"
                    comision_pct = 0.06
                elif pct >= 0.80:
                    estatus = "REGULAR"
                    comision_pct = 0.02
                else:
                    estatus = "DEFICIENTE"
                    comision_pct = 0.0

                comision_var = round(r["LOGRO_REAL"] * comision_pct, 2)
                bono_fijo = 3000.0 if pct >= 1.0 else 0.0
                comp_total = round(comision_var + bono_fijo, 2)

                c_ws_task.append([r["ID_EMPLEADO"], r["NOMBRE"], r["SUCURSAL"], r["META_ASIGNADA"], r["LOGRO_REAL"], "", "", "", "", ""])
                s_ws_task.append([
                    r["ID_EMPLEADO"],
                    r["NOMBRE"],
                    r["SUCURSAL"],
                    r["META_ASIGNADA"],
                    r["LOGRO_REAL"],
                    pct,
                    estatus,
                    comision_var,
                    bono_fijo,
                    comp_total,
                ])
                s_ws_task.cell(row=idx, column=4).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=5).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "0.0%"
                s_ws_task.cell(row=idx, column=8).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=9).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=10).number_format = "$#,##0.00"

                expected_summary[f"G{idx}"] = estatus
                expected_summary[f"H{idx}"] = comision_var
                expected_summary[f"J{idx}"] = comp_total
                criteria_status.append(f"G{idx}")
                criteria_bonus.extend([f"H{idx}", f"J{idx}"])

            else:  # intermediate
                estatus = "CUMPLIDO" if pct >= 1.0 else "REGULAR" if pct >= 0.80 else "DEFICIENTE"
                bono = 5000.0 if estatus == "CUMPLIDO" else 2000.0 if estatus == "REGULAR" else 0.0

                c_ws_task.append([r["ID_EMPLEADO"], r["NOMBRE"], r["SUCURSAL"], r["META_ASIGNADA"], r["LOGRO_REAL"], "", "", ""])
                s_ws_task.append([
                    r["ID_EMPLEADO"],
                    r["NOMBRE"],
                    r["SUCURSAL"],
                    r["META_ASIGNADA"],
                    r["LOGRO_REAL"],
                    pct,
                    estatus,
                    bono,
                ])
                s_ws_task.cell(row=idx, column=4).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=5).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "0.0%"
                s_ws_task.cell(row=idx, column=8).number_format = "$#,##0.00"

                expected_summary[f"G{idx}"] = estatus
                expected_summary[f"H{idx}"] = bono
                criteria_status.append(f"G{idx}")
                criteria_bonus.append(f"H{idx}")

            for c in range(1, num_cols + 1):
                c_ws_task.cell(row=idx, column=c).border = cls.THIN_BORDER
                s_ws_task.cell(row=idx, column=c).border = cls.THIN_BORDER

        tot_row = len(rows) + 12
        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=tot_row, column=1, value="TOTAL COMPENSACIONES:").font = cls.BOLD_FONT
            for col in range(1, num_cols + 1):
                ws.cell(row=tot_row, column=col).border = cls.TOTAL_BORDER

        if difficulty == "advanced":
            tot_comp = round(sum(
                round(r["LOGRO_REAL"] * (0.10 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 1.15 else 0.06 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 1.0 else 0.02 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 0.8 else 0.0), 2)
                + (3000.0 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 1.0 else 0.0)
                for r in rows
            ), 2)
            s_ws_task.cell(row=tot_row, column=10, value=tot_comp).number_format = "$#,##0.00"
            expected_summary[f"J{tot_row}"] = tot_comp
            criteria_bonus.append(f"J{tot_row}")
        else:
            tot_bonos = sum(
                (3000.0 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 1.0 else 0.0) if difficulty == "basic"
                else (5000.0 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 1.0 else 2000.0 if r["LOGRO_REAL"]/r["META_ASIGNADA"] >= 0.8 else 0.0)
                for r in rows
            )
            s_ws_task.cell(row=tot_row, column=8, value=tot_bonos).number_format = "$#,##0.00"
            expected_summary[f"H{tot_row}"] = tot_bonos
            criteria_bonus.append(f"H{tot_row}")

        expected_summary["__criteria__"] = {
            "Condicionales de Estatus": criteria_status,
            "Cálculo de Bonificaciones": criteria_bonus,
        }

        crit_ws = s_wb.create_sheet(title="Criterios")
        crit_ws.append(["Criterio", "Rango_Celdas"])
        crit_ws.append(["Condicionales de Estatus", f"G12:G{tot_row-1}"])
        crit_ws.append(["Cálculo de Bonificaciones", f"H12:J{tot_row}"])

        cls._auto_fit_columns(c_ws_task)
        cls._auto_fit_columns(s_ws_task)

        return cls._save_to_bytes(c_wb), cls._save_to_bytes(s_wb), expected_summary

    # ---------------------------------------------------------
    # 4. BLUEPRINT: Conciliación y Auditoría
    # ---------------------------------------------------------
    @classmethod
    def _build_audit_reconciliation(
        cls, bp: Any, title: str, data: dict[str, Any], difficulty: str = "advanced"
    ) -> tuple[bytes, bytes, dict[str, Any]]:
        c_wb = openpyxl.Workbook()
        c_ws_task = c_wb.active
        c_ws_task.title = bp.task_sheet_name
        c_ws_src = c_wb.create_sheet(title=bp.source_sheet_name)
        
        s_wb = openpyxl.Workbook()
        s_ws_task = s_wb.active
        s_ws_task.title = bp.task_sheet_name
        s_ws_src = s_wb.create_sheet(title=bp.source_sheet_name)

        rows = data["rows"]

        # 1. Fill DatosAuditoria in both
        src_headers = ["CODIGO", "DESCRIPCION", "COSTO_UNITARIO", "REGISTRO_SISTEMA", "CONTEO_FISICO"]
        for ws in (c_ws_src, s_ws_src):
            ws.append(src_headers)
            for c in range(1, len(src_headers) + 1):
                ws.cell(row=1, column=c).fill = cls.HEADER_FILL
                ws.cell(row=1, column=c).font = cls.HEADER_FONT
                ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center")
            for r in rows:
                ws.append([r["CODIGO"], r["DESCRIPCION"], r["COSTO_UNITARIO"], r["REGISTRO_SISTEMA"], r["CONTEO_FISICO"]])
                ws.cell(row=ws.max_row, column=3).number_format = "$#,##0.00"
                ws.cell(row=ws.max_row, column=4).number_format = "#,##0"
                ws.cell(row=ws.max_row, column=5).number_format = "#,##0"
            cls._auto_fit_columns(ws)

        # 2. Instructions
        inst_text = get_instructions_for_blueprint(bp.id, difficulty)
        for ws in (c_ws_task, s_ws_task):
            cls._create_instruction_banner(ws, f"{title} ({difficulty.upper()})", inst_text)

        # 3. Task Headers based on difficulty
        if difficulty == "basic":
            headers = ["CODIGO", "DESCRIPCION", "COSTO_UNITARIO", "REGISTRO_SISTEMA", "CONTEO_FISICO", "DIFERENCIA_UNIDADES (Fórmula)", "IMPORTE_DISCREPANCIA (Fórmula)"]
        elif difficulty == "advanced":
            headers = ["CODIGO", "DESCRIPCION", "COSTO_UNITARIO", "REGISTRO_SISTEMA", "CONTEO_FISICO", "DIFERENCIA_UNIDADES (Fórmula)", "IMPORTE_DISCREPANCIA (Fórmula)", "% DESVIACION (Fórmula)", "CRITICIDAD (Fórmula)", "DICTAMEN (Fórmula)"]
        else:  # intermediate
            headers = ["CODIGO", "DESCRIPCION", "COSTO_UNITARIO", "REGISTRO_SISTEMA", "CONTEO_FISICO", "DIFERENCIA_UNIDADES (Fórmula)", "IMPORTE_DISCREPANCIA (Fórmula)", "DICTAMEN (Fórmula)"]

        num_cols = len(headers)

        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=10, column=1, value="MATRIZ DE CONCILIACIÓN Y DICTAMEN DE AUDITORÍA:").font = cls.BOLD_FONT
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=11, column=col_idx, value=h)
                cell.fill = cls.HEADER_FILL
                cell.font = cls.HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")

        expected_summary = {}
        criteria_diff = []
        criteria_dictamen = []

        for idx, r in enumerate(rows, start=12):
            diff_u = r["CONTEO_FISICO"] - r["REGISTRO_SISTEMA"]
            diff_amt = round(diff_u * r["COSTO_UNITARIO"], 2)
            dictamen = "CUADRADO" if diff_u == 0 else "SOBRANTE" if diff_u > 0 else "FALTANTE"

            if difficulty == "basic":
                c_ws_task.append([r["CODIGO"], r["DESCRIPCION"], r["COSTO_UNITARIO"], r["REGISTRO_SISTEMA"], r["CONTEO_FISICO"], "", ""])
                s_ws_task.append([
                    r["CODIGO"],
                    r["DESCRIPCION"],
                    r["COSTO_UNITARIO"],
                    r["REGISTRO_SISTEMA"],
                    r["CONTEO_FISICO"],
                    diff_u,
                    diff_amt,
                ])
                s_ws_task.cell(row=idx, column=3).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "#,##0"
                s_ws_task.cell(row=idx, column=7).number_format = "$#,##0.00"

                expected_summary[f"F{idx}"] = diff_u
                expected_summary[f"G{idx}"] = diff_amt
                criteria_diff.extend([f"F{idx}", f"G{idx}"])

            elif difficulty == "advanced":
                pct_desv = round(diff_u / max(r["REGISTRO_SISTEMA"], 1), 4)
                criticidad = "CRÍTICO" if (abs(diff_amt) >= 5000 or abs(pct_desv) >= 0.15) else "TOLERABLE"

                c_ws_task.append([r["CODIGO"], r["DESCRIPCION"], r["COSTO_UNITARIO"], r["REGISTRO_SISTEMA"], r["CONTEO_FISICO"], "", "", "", "", ""])
                s_ws_task.append([
                    r["CODIGO"],
                    r["DESCRIPCION"],
                    r["COSTO_UNITARIO"],
                    r["REGISTRO_SISTEMA"],
                    r["CONTEO_FISICO"],
                    diff_u,
                    diff_amt,
                    pct_desv,
                    criticidad,
                    dictamen,
                ])
                s_ws_task.cell(row=idx, column=3).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "#,##0"
                s_ws_task.cell(row=idx, column=7).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=8).number_format = "0.0%"

                expected_summary[f"F{idx}"] = diff_u
                expected_summary[f"G{idx}"] = diff_amt
                expected_summary[f"I{idx}"] = criticidad
                expected_summary[f"J{idx}"] = dictamen
                criteria_diff.extend([f"F{idx}", f"G{idx}"])
                criteria_dictamen.extend([f"I{idx}", f"J{idx}"])

            else:  # intermediate
                c_ws_task.append([r["CODIGO"], r["DESCRIPCION"], r["COSTO_UNITARIO"], r["REGISTRO_SISTEMA"], r["CONTEO_FISICO"], "", "", ""])
                s_ws_task.append([
                    r["CODIGO"],
                    r["DESCRIPCION"],
                    r["COSTO_UNITARIO"],
                    r["REGISTRO_SISTEMA"],
                    r["CONTEO_FISICO"],
                    diff_u,
                    diff_amt,
                    dictamen,
                ])
                s_ws_task.cell(row=idx, column=3).number_format = "$#,##0.00"
                s_ws_task.cell(row=idx, column=6).number_format = "#,##0"
                s_ws_task.cell(row=idx, column=7).number_format = "$#,##0.00"

                expected_summary[f"F{idx}"] = diff_u
                expected_summary[f"G{idx}"] = diff_amt
                expected_summary[f"H{idx}"] = dictamen
                criteria_diff.extend([f"F{idx}", f"G{idx}"])
                criteria_dictamen.append(f"H{idx}")

            for c in range(1, num_cols + 1):
                c_ws_task.cell(row=idx, column=c).border = cls.THIN_BORDER
                s_ws_task.cell(row=idx, column=c).border = cls.THIN_BORDER

        tot_row = len(rows) + 12
        tot_diff_amount = round(sum(round((r["CONTEO_FISICO"] - r["REGISTRO_SISTEMA"]) * r["COSTO_UNITARIO"], 2) for r in rows), 2)

        for ws in (c_ws_task, s_ws_task):
            ws.cell(row=tot_row, column=1, value="IMPORTE NETO DE DISCREPANCIA:").font = cls.BOLD_FONT
            for col in range(1, num_cols + 1):
                ws.cell(row=tot_row, column=col).border = cls.TOTAL_BORDER

        s_ws_task.cell(row=tot_row, column=7, value=tot_diff_amount).number_format = "$#,##0.00"
        expected_summary[f"G{tot_row}"] = tot_diff_amount
        criteria_diff.append(f"G{tot_row}")

        expected_summary["__criteria__"] = {
            "Diferencias Físicas y Monetarias": criteria_diff,
            "Dictamen Condicional": criteria_dictamen,
        }

        crit_ws = s_wb.create_sheet(title="Criterios")
        crit_ws.append(["Criterio", "Rango_Celdas"])
        crit_ws.append(["Diferencias Físicas y Monetarias", f"F12:G{tot_row}"])
        if criteria_dictamen:
            crit_ws.append(["Dictamen Condicional", f"H12:J{tot_row-1}"])

        cls._auto_fit_columns(c_ws_task)
        cls._auto_fit_columns(s_ws_task)

        return cls._save_to_bytes(c_wb), cls._save_to_bytes(s_wb), expected_summary

    # ---------------------------------------------------------
    # Helper Utilities
    # ---------------------------------------------------------
    @classmethod
    def _create_instruction_banner(cls, ws: Any, title: str, instructions: str) -> None:
        ws.merge_cells("A1:I2")
        banner = ws["A1"]
        banner.value = f"PRUEBA TÉCNICA PRÁCTICA: {title.upper()}"
        banner.fill = cls.BANNER_FILL
        banner.font = cls.BANNER_FONT
        banner.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A4:I7")
        box = ws["A4"]
        box.value = instructions
        box.fill = cls.CARD_FILL
        box.font = cls.DATA_FONT
        box.alignment = Alignment(vertical="top", wrap_text=True)
        
        for r in range(4, 8):
            for c in range(1, 10):
                ws.cell(row=r, column=c).border = cls.CARD_BORDER

    @classmethod
    def _auto_fit_columns(cls, ws: Any) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2, 4, 5, 6, 7]:  # Skip banner and instructions rows
                    continue
                val_str = str(cell.value or "")
                if val_str:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def _save_to_bytes(cls, wb: openpyxl.Workbook) -> bytes:
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
