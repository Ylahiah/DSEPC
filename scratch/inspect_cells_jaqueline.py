import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

path_sub = r'C:\Users\MSI\Desktop\Python\DSEPC\backend\storage\excel_submissions\48\581_45bf7856743a4f26b7e82b5bce94d103.xlsx'
wb = openpyxl.load_workbook(path_sub, data_only=True)
ws = wb['Evaluación Tablas Dinámicas']

print("--- ROWS 18-26 (Ejercicio 1: Meses) ---")
for r in range(18, 27):
    vals = [ws.cell(r, c).value for c in range(1, 8)]
    print(f"Row {r:2d}: {vals}")

print("\n--- ROWS 47-57 (Ejercicio 2: Jurisdicciones) ---")
for r in range(47, 58):
    vals = [ws.cell(r, c).value for c in range(1, 8)]
    print(f"Row {r:2d}: {vals}")

print("\n--- ROWS 75-89 (Ejercicio 3: Capturistas) ---")
for r in range(75, 90):
    vals = [ws.cell(r, c).value for c in range(1, 8)]
    print(f"Row {r:2d}: {vals}")
